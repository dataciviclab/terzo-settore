#!/usr/bin/env python3
"""Parse ISTAT Excel 2023 into a single combined CSV for the pipeline."""
import csv
import sys
from pathlib import Path

def parse_istat_non_profit(excel_path: str, output_path: str):
    """Parse ISTAT Excel 2023 non-profit tables into a single long-format CSV."""
    try:
        import openpyxl
    except ImportError:
        print("ERROR: openpyxl not installed", file=sys.stderr)
        sys.exit(1)

    wb = openpyxl.load_workbook(excel_path, read_only=True)
    rows_out = []

    # --- Table 1: Regioni per forma giuridica ---
    ws1 = wb['1']
    rows1 = list(ws1.iter_rows(values_only=True))
    fg_names = ['Associazione', 'Cooperativa sociale', 'Fondazione', 'Altra forma giuridica']

    for row in rows1[5:]:
        name = str(row[0] or '').strip()
        if not name or any(name.startswith(p) for p in ('Nord', 'Centro', 'Sud', 'Isole')):
            continue
        for j, fg in enumerate(fg_names):
            ist = row[j + 1]
            dip = row[j + 7] if len(row) > j + 7 else None
            if ist is not None:
                rows_out.append({
                    'anno': 2023, 'regione': name, 'livello': 'regione',
                    'categoria_tipo': 'forma_giuridica', 'categoria': fg,
                    'istituzioni': int(ist), 'dipendenti': int(dip) if dip else 0,
                })

    # --- Table 2: Regioni per settore ---
    ws2 = wb['2']
    rows2 = list(ws2.iter_rows(values_only=True))
    raw_sectors = [str(v or '').strip() for v in rows2[3][1:]]
    # First14 = ist, skip Totale, next14 = dip
    sectors = [s for s in raw_sectors[:15] if s and s != 'Totale']

    for row in rows2[5:]:
        name = str(row[0] or '').strip()
        if not name or any(name.startswith(p) for p in ('Nord', 'Centro', 'Sud', 'Isole')):
            continue
        for j, s in enumerate(sectors):
            ist = row[j + 1]
            dip = row[j + 16] if len(row) > j + 16 else None
            if ist is not None:
                rows_out.append({
                    'anno': 2023, 'regione': name, 'livello': 'regione',
                    'categoria_tipo': 'settore', 'categoria': s,
                    'istituzioni': int(ist), 'dipendenti': int(dip) if dip else 0,
                })

    # --- Table 22: Province ---
    ws22 = wb['22']
    rows22 = list(ws22.iter_rows(values_only=True))
    for row in rows22[2:]:
        code = str(row[0] or '').strip()
        prov = str(row[1] or '').strip()
        ist = row[2]
        dip = row[3]
        if code and prov and ist is not None:
            rows_out.append({
                'anno': 2023, 'regione': prov, 'livello': 'provincia',
                'categoria_tipo': 'totale', 'categoria': 'totale',
                'istituzioni': int(ist), 'dipendenti': int(dip) if dip else 0,
                'codice_provincia': code,
            })

    wb.close()

    fieldnames = ['anno', 'regione', 'livello', 'categoria_tipo', 'categoria',
                  'istituzioni', 'dipendenti', 'codice_provincia']
    with open(output_path, 'w', newline='', encoding='utf-8') as f:
        w = csv.DictWriter(f, fieldnames=fieldnames, extrasaction='ignore')
        w.writeheader()
        w.writerows(rows_out)

    print(f"Parsed {len(rows_out)} rows -> {output_path}")


if __name__ == '__main__':
    excel = sys.argv[1] if len(sys.argv) > 1 else '/tmp/tavole_nonprofit_2023.xlsx'
    output = sys.argv[2] if len(sys.argv) > 2 else 'data/istat_non_profit_2023.csv'
    parse_istat_non_profit(excel, output)
